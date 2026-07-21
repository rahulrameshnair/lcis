"""Template-driven LCIS metadata collection and export.

This file owns the LCIS metadata template workflow. It reads the packaged Excel
template, converts template rows into structured metadata fields, prompts users
for values, fills metadata workbooks, and writes JSON/CSV sidecars. It depends
on openpyxl at runtime for Excel workbooks, click for interactive prompts,
csv/json for machine-readable outputs, and importlib.resources for locating the
packaged template.
"""

from __future__ import annotations

import csv
import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from importlib.resources import files
from pathlib import Path
from typing import Any


STRUCTURE_SHEET = "_structure"
METADATA_SHEET = "Metadata"
LICENSES_SHEET = "_Infolicenses"
DEFAULT_TEMPLATE = "metadata_template.xlsx"


@dataclass(frozen=True)
class MetadataField:
    """Description of one metadata field defined by the template workbook."""
    """
    Each instance is a compact record of one row from the ``_structure`` tab.
    The ``structure_row`` value is kept because the visible Metadata sheet can
    refer back to structure rows through Excel formulas, so later read/write
    operations can identify fields even when labels are reused or moved.
    """

    element: str
    attribute: str
    label: str
    definition: str
    datatype: str
    guidance: str
    structure_row: int


def default_template_path() -> Path:
    """Return the packaged metadata template path."""
    """
    The template is shipped inside ``lcis.templates`` rather than discovered
    from the current working directory. This keeps command-line metadata
    creation reproducible regardless of where the user runs the command.
    """

    return Path(files("lcis.templates") / DEFAULT_TEMPLATE)


def load_metadata_fields(template_path: str | Path | None = None) -> list[MetadataField]:
    """Read metadata fields from the template's ``_structure`` sheet."""
    """
    The ``_structure`` sheet is treated as the authoritative schema for the
    metadata wizard. Rows can inherit the current element from previous rows,
    while every row with an attribute becomes a ``MetadataField`` that drives
    prompting, workbook filling, and CSV ordering.
    """

    from openpyxl import load_workbook

    template_path = Path(template_path) if template_path else default_template_path()
    workbook = load_workbook(template_path, data_only=False, read_only=True)
    worksheet = workbook[STRUCTURE_SHEET]

    fields: list[MetadataField] = []
    current_element: str | None = None

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, max_col=6),
        start=2,
    ):
        element, attribute, label, definition, datatype, guidance = [
            cell.value for cell in row[:6]
        ]
        if element:
            current_element = str(element)
        if not current_element or not attribute:
            continue

        fields.append(
            MetadataField(
                element=current_element,
                attribute=str(attribute),
                label=str(label or attribute),
                definition=str(definition or ""),
                datatype=str(datatype or "Text"),
                guidance=str(guidance or ""),
                structure_row=row_number,
            )
        )

    workbook.close()
    return fields


def load_license_lookup(template_path: str | Path | None = None) -> dict[str, str]:
    """Read license name to SPDX identifier mappings from the template."""
    """
    The license helper keeps common SPDX identifiers inside the same template
    ecosystem as the metadata fields. When a user chooses or enters a known
    license name, LCIS can fill the matching SPDX value without forcing the
    user to remember the identifier manually.
    """

    from openpyxl import load_workbook

    template_path = Path(template_path) if template_path else default_template_path()
    workbook = load_workbook(template_path, data_only=True, read_only=True)
    worksheet = workbook[LICENSES_SHEET]
    lookup: dict[str, str] = {}

    for name, spdx in worksheet.iter_rows(min_row=1, max_col=2, values_only=True):
        if name:
            lookup[str(name)] = "" if spdx is None else str(spdx)

    workbook.close()
    return lookup


def collect_metadata_from_template(
    *,
    template_path: str | Path | None = None,
    initial_values: dict[tuple[str, str], Any] | None = None,
) -> dict[str, dict[str, Any]]:
    """Prompt for metadata values using fields from the Excel template."""
    """
    This is the interactive heart of the metadata wizard. It walks through the
    template fields in workbook order, groups prompts by LCIS element, stores
    blank answers as ``None``, adds an explicit visual break after each answer
    for notebook frontends, enriches License.spdx when possible, and adds a
    timestamp under ``_meta`` so generated sidecar files can be traced.
    """

    import click

    fields = load_metadata_fields(template_path)
    initial_values = initial_values or {}
    metadata: dict[str, dict[str, Any]] = {}
    current_element: str | None = None

    click.echo("")
    click.secho("  LCIS2024 Metadata Wizard", fg="cyan", bold=True)
    click.secho("  " + "-" * 40, fg="cyan")
    click.echo(
        "  Values are collected from the metadata Excel template.\n"
        "  Leave optional fields blank when they do not apply.\n"
    )

    for field in fields:
        if field.element != current_element:
            current_element = field.element
            click.echo("")
            click.secho(f"  {current_element}", fg="cyan", bold=True)

        metadata.setdefault(field.element, {})

        default = initial_values.get((field.element, field.attribute), "")
        if field.definition:
            click.secho(f"  {field.label}: {field.definition}", fg="bright_black")
        if field.guidance:
            click.secho(f"  Format: {field.guidance}", fg="bright_black")

        raw = click.prompt(
            f"  {field.label}",
            default="" if default is None else str(default),
            show_default=bool(default),
        )
        value = raw.strip()
        metadata[field.element][field.attribute] = value or None
        click.echo("")

    _fill_license_spdx(metadata, template_path)
    metadata["_meta"] = {
        "created_at": datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    }
    return metadata


def _fill_license_spdx(
    metadata: dict[str, dict[str, Any]],
    template_path: str | Path | None = None,
) -> None:
    """Fill ``License.spdx`` from the template lookup when it is missing."""
    """
    This helper is deliberately conservative: it only writes an SPDX value
    when the metadata already has a license name and no SPDX value. That keeps
    manually entered identifiers intact while still helping template-driven
    workflows produce more complete machine-readable metadata.
    """

    license_data = metadata.get("License")
    if not license_data or license_data.get("spdx"):
        return

    license_name = license_data.get("name")
    if not license_name:
        return

    lookup = load_license_lookup(template_path)
    spdx = lookup.get(str(license_name))
    if spdx:
        license_data["spdx"] = spdx


def write_metadata_json(metadata: dict[str, Any], output_path: str | Path) -> Path:
    """Write metadata as formatted UTF-8 JSON and return the output path."""
    """
    JSON is the machine-readable sidecar used by downstream tools. The writer
    creates missing parent directories, preserves non-ASCII metadata text, and
    keeps indentation stable so researchers can inspect or diff the file.
    """

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as file:
        json.dump(metadata, file, indent=2, ensure_ascii=False)
    return output_path


def write_metadata_csv(
    metadata: dict[str, dict[str, Any]],
    output_path: str | Path,
    *,
    template_path: str | Path | None = None,
) -> Path:
    """Write template-ordered metadata values to a simple CSV file."""
    """
    The CSV export is organized by the same field order as the Excel template
    instead of by dictionary insertion order. This keeps the spreadsheet,
    prompt sequence, and CSV sidecar aligned for review and documentation.
    """

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fields = load_metadata_fields(template_path)

    with open(output_path, "w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(["Element", "Attribute", "Label", "Value"])
        for field in fields:
            value = metadata.get(field.element, {}).get(field.attribute)
            writer.writerow([field.element, field.attribute, field.label, value or ""])

    return output_path


def write_metadata_workbook(
    metadata: dict[str, dict[str, Any]],
    output_path: str | Path,
    *,
    template_path: str | Path | None = None,
) -> Path:
    """Copy the metadata template and fill the human-facing ``Metadata`` tab."""
    """
    The workbook writer starts from a clean copy of the template, then maps
    visible rows back to ``MetadataField`` records. It supports both formula
    references to the structure sheet and plain label matching, which makes it
    tolerant of small template layout changes while keeping values in column 3.
    """

    from openpyxl import load_workbook

    template_path = Path(template_path) if template_path else default_template_path()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    workbook = load_workbook(output_path)
    metadata_sheet = workbook[METADATA_SHEET]
    field_by_row = {field.structure_row: field for field in load_metadata_fields(template_path)}
    fields = load_metadata_fields(template_path)

    current_element: str | None = None
    for row_number in range(1, metadata_sheet.max_row + 1):
        element_cell = metadata_sheet.cell(row=row_number, column=1).value
        label_cell = metadata_sheet.cell(row=row_number, column=2).value

        row_field = _field_from_formula(label_cell, field_by_row)
        if row_field is None:
            current_element = _element_from_cell(element_cell, current_element, field_by_row)
            row_field = _field_from_label(current_element, label_cell, fields)
        else:
            current_element = row_field.element

        if row_field is None:
            current_element = _element_from_cell(element_cell, current_element, field_by_row)
            continue

        value = metadata.get(row_field.element, {}).get(row_field.attribute)
        metadata_sheet.cell(row=row_number, column=3).value = value or None

    workbook.save(output_path)
    workbook.close()
    return output_path


def read_metadata_workbook(
    workbook_path: str | Path,
    *,
    template_path: str | Path | None = None,
) -> dict[str, dict[str, Any]]:
    """Read filled values from the human-facing ``Metadata`` tab."""
    """
    This function performs the inverse of ``write_metadata_workbook``. It
    walks the visible sheet, identifies each row's LCIS element/attribute, and
    reconstructs the nested metadata dictionary so a manually filled workbook
    can be converted into JSON and CSV without re-running the wizard.
    """

    from openpyxl import load_workbook

    workbook_path = Path(workbook_path)
    template_path = Path(template_path) if template_path else workbook_path
    workbook = load_workbook(workbook_path, data_only=False)
    metadata_sheet = workbook[METADATA_SHEET]
    field_by_row = {field.structure_row: field for field in load_metadata_fields(template_path)}
    fields = load_metadata_fields(template_path)

    metadata: dict[str, dict[str, Any]] = {}
    current_element: str | None = None

    for row_number in range(1, metadata_sheet.max_row + 1):
        element_cell = metadata_sheet.cell(row=row_number, column=1).value
        label_cell = metadata_sheet.cell(row=row_number, column=2).value
        value = metadata_sheet.cell(row=row_number, column=3).value

        row_field = _field_from_formula(label_cell, field_by_row)
        if row_field is None:
            current_element = _element_from_cell(element_cell, current_element, field_by_row)
            row_field = _field_from_label(current_element, label_cell, fields)
        else:
            current_element = row_field.element

        if row_field is None:
            current_element = _element_from_cell(element_cell, current_element, field_by_row)
            continue

        metadata.setdefault(row_field.element, {})
        metadata[row_field.element][row_field.attribute] = value

    workbook.close()
    _fill_license_spdx(metadata, template_path)
    metadata["_meta"] = {
        "created_at": datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    }
    return metadata


def write_all_metadata_outputs(
    metadata: dict[str, dict[str, Any]],
    output_stem: str,
    output_dir: str | Path = ".",
    *,
    template_path: str | Path | None = None,
) -> dict[str, Path]:
    """Write XLSX, JSON, and CSV metadata outputs with a shared stem."""
    """
    This is the synchronization point for metadata sidecars. Callers provide
    one nested metadata dictionary, and this helper writes all supported output
    formats together so the workbook, JSON, and CSV represent the same answers.
    """

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    xlsx_path = write_metadata_workbook(
        metadata,
        output_dir / f"{output_stem}_metadata.xlsx",
        template_path=template_path,
    )
    json_path = write_metadata_json(metadata, output_dir / f"{output_stem}_metadata.json")
    csv_path = write_metadata_csv(
        metadata,
        output_dir / f"{output_stem}_metadata.csv",
        template_path=template_path,
    )

    return {"xlsx": xlsx_path, "json": json_path, "csv": csv_path}


def convert_metadata_workbook(
    workbook_path: str | Path,
    *,
    output_stem: str | None = None,
    output_dir: str | Path | None = None,
) -> dict[str, Path]:
    """Convert a filled metadata workbook into JSON and CSV sidecar files.

    This keeps the workbook untouched. It is intended for users who prefer to
    fill the human-facing Excel metadata sheet manually and then generate the
    machine-readable files afterward.
    """
    """
    The converter supports the manual Excel-first workflow. It reads the
    workbook as the source of truth, derives a sensible output stem when the
    caller does not provide one, and writes only the sidecar formats so the
    original filled workbook remains unchanged.
    """

    workbook_path = Path(workbook_path)
    output_dir = Path(output_dir) if output_dir else workbook_path.parent
    output_stem = output_stem or _metadata_stem_from_workbook(workbook_path)

    metadata = read_metadata_workbook(workbook_path)
    json_path = write_metadata_json(
        metadata,
        output_dir / f"{output_stem}_metadata.json",
    )
    csv_path = write_metadata_csv(
        metadata,
        output_dir / f"{output_stem}_metadata.csv",
        template_path=workbook_path,
    )

    return {"json": json_path, "csv": csv_path}


def _field_from_formula(
    value: Any,
    field_by_row: dict[int, MetadataField],
) -> MetadataField | None:
    """Resolve a metadata field from a formula pointing at ``_structure``."""
    """
    Some rows in the visible Metadata sheet are linked to the hidden structure
    sheet through formulas. This parser extracts the referenced row number and
    uses it to recover the exact template field behind that visible label.
    """

    if not isinstance(value, str) or not value.startswith("="):
        return None
    match = re.search(r"_structure!\$?[A-Z]+\$?(\d+)", value, re.IGNORECASE)
    if not match:
        return None
    return field_by_row.get(int(match.group(1)))


def _metadata_stem_from_workbook(workbook_path: Path) -> str:
    """Derive an output stem from a metadata workbook filename."""
    """
    Filled metadata workbooks normally end in ``_metadata.xlsx``. Removing
    that suffix keeps generated JSON/CSV sidecars named consistently with the
    dataset rather than repeating the metadata suffix twice.
    """

    stem = workbook_path.stem
    suffix = "_metadata"
    if stem.endswith(suffix):
        return stem[: -len(suffix)]
    return stem


def _element_from_cell(
    value: Any,
    current_element: str | None,
    field_by_row: dict[int, MetadataField],
) -> str | None:
    """Track the current metadata element while scanning workbook rows."""
    """
    The visible sheet may show element names directly or through formulas.
    This helper updates the current element when a row declares one and keeps
    the previous element for rows where the template leaves that cell blank.
    """

    if isinstance(value, str) and value.startswith("="):
        field = _field_from_formula(value, field_by_row)
        return field.element if field else current_element
    if value:
        return str(value)
    return current_element


def _field_from_label(
    element: str | None,
    label: Any,
    fields: list[MetadataField],
) -> MetadataField | None:
    """Find a field by element and visible label when no formula is present."""
    """
    This fallback lets LCIS read workbooks where the visible Metadata sheet is
    not formula-linked to ``_structure``. It is less strict than row matching,
    but useful for manually adjusted templates that preserve labels.
    """

    if not element or not label:
        return None
    label_text = str(label)
    for field in fields:
        if field.element == element and field.label == label_text:
            return field
    return None


__all__ = [
    "MetadataField",
    "collect_metadata_from_template",
    "convert_metadata_workbook",
    "default_template_path",
    "load_license_lookup",
    "load_metadata_fields",
    "read_metadata_workbook",
    "write_all_metadata_outputs",
    "write_metadata_csv",
    "write_metadata_json",
    "write_metadata_workbook",
]
